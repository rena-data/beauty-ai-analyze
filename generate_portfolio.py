"""취업용 포트폴리오 PPT - 모던 프로페셔널 디자인"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

SS = "/Users/data_team/Desktop/rena_zip/project/beauty_ai_analyze/screenshots"
OUT = "/Users/data_team/Desktop/rena_zip/project/beauty_ai_analyze"

# ─── Design System ───
W = RGBColor(0xFF, 0xFF, 0xFF)
BG = RGBColor(0xF5, 0xF5, 0xF7)
BLK = RGBColor(0x1A, 0x1A, 0x2E)
DK = RGBColor(0x16, 0x21, 0x3E)
GR = RGBColor(0x8B, 0x8B, 0x8B)
LG = RGBColor(0xE0, 0xE0, 0xE0)
GN = RGBColor(0x00, 0xC9, 0xA7)
BL = RGBColor(0x38, 0x7A, 0xFF)
RD = RGBColor(0xFF, 0x5C, 0x5C)


def nprs():
    p = Presentation(); p.slide_width = Inches(13.333); p.slide_height = Inches(7.5); return p

def sbg(s, c=BG):
    f = s.background.fill; f.solid(); f.fore_color.rgb = c

def rect(s, l, t, w, h, fill=W, bd=None):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid(); sh.fill.fore_color.rgb = fill; sh.line.fill.background()
    return sh

def rrect(s, l, t, w, h, fill=W):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid(); sh.fill.fore_color.rgb = fill; sh.line.fill.background()
    return sh

def tx(s, l, t, w, h, text, sz=14, b=False, c=BLK, a=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(l, t, w, h); tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text; p.font.size = Pt(sz); p.font.bold = b; p.font.color.rgb = c; p.alignment = a

def bl(s, l, t, w, h, items, sz=11, c=BLK, sp=3):
    tb = s.shapes.add_textbox(l, t, w, h); tf = tb.text_frame; tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item; p.font.size = Pt(sz); p.font.color.rgb = c; p.space_after = Pt(sp)

def im(s, path, l, t, w=None, h=None):
    if not os.path.exists(path): return
    if w: s.shapes.add_picture(path, l, t, width=w)
    elif h: s.shapes.add_picture(path, l, t, height=h)

def sidebar(s, accent, width=Inches(0.25)):
    rect(s, 0, 0, width, Inches(7.5), fill=accent)

def tag(s, l, t, text, accent):
    sh = rrect(s, l, t, Inches(1.2), Inches(0.3), fill=accent)
    tx(s, l, t + Inches(0.02), Inches(1.2), Inches(0.25), text, 9, True, W, PP_ALIGN.CENTER)


# ═══ COVER ═══
def cover(prs, title, sub, accent):
    s = prs.slides.add_slide(prs.slide_layouts[6]); sbg(s, DK)
    # 좌측 액센트 바
    rect(s, 0, 0, Inches(0.4), Inches(7.5), fill=accent)
    # 중앙 콘텐츠
    tx(s, Inches(1.5), Inches(1.5), Inches(10), Inches(0.4), "PORTFOLIO", 14, True, accent, PP_ALIGN.LEFT)
    tx(s, Inches(1.5), Inches(2.2), Inches(10), Inches(1), title, 44, True, W)
    tx(s, Inches(1.5), Inches(3.5), Inches(10), Inches(0.6), sub, 16, c=GR)
    # 하단 정보
    rect(s, Inches(1.5), Inches(5.5), Inches(10), Inches(0.01), fill=accent)
    tx(s, Inches(1.5), Inches(5.8), Inches(3), Inches(0.4), "이기쁨  Lee Gippeum", 14, True, W)
    tx(s, Inches(1.5), Inches(6.3), Inches(5), Inches(0.3), "ldsjoy@naver.com  |  github.com/rena-data", 11, c=GR)
    tx(s, Inches(10), Inches(5.8), Inches(2.5), Inches(0.4), "2026.05", 14, c=GR, a=PP_ALIGN.RIGHT)


# ═══ ABOUT ME ═══
def about(prs, accent, version):
    s = prs.slides.add_slide(prs.slide_layouts[6]); sbg(s, BG)
    sidebar(s, accent)

    # 좌측 프로필
    rrect(s, Inches(0.8), Inches(0.5), Inches(4.2), Inches(6.5), fill=DK)
    tx(s, Inches(1.2), Inches(0.8), Inches(3.5), Inches(0.5), "이기쁨", 28, True, W)
    tx(s, Inches(1.2), Inches(1.4), Inches(3.5), Inches(0.3), "Lee Gippeum", 12, c=GR)

    labels = {"data": "데이터 분석가", "plan": "서비스 기획자", "dev": "백엔드 개발자"}
    tag(s, Inches(1.2), Inches(1.9), labels[version], accent)

    tx(s, Inches(1.2), Inches(2.5), Inches(3.5), Inches(0.3), "Contact", 10, True, accent)
    bl(s, Inches(1.2), Inches(2.9), Inches(3.5), Inches(0.8), [
        "ldsjoy@naver.com",
        "github.com/rena-data",
    ], sz=10, c=RGBColor(0xCC, 0xCC, 0xCC))

    tx(s, Inches(1.2), Inches(3.8), Inches(3.5), Inches(0.3), "Education", 10, True, accent)
    bl(s, Inches(1.2), Inches(4.2), Inches(3.5), Inches(0.8), [
        "스파르타코딩클럽 데이터 분석 부트캠프",
        "2024.11 ~ 2025.03",
        "",
        "웅지세무대학교 부동산금융평가과",
    ], sz=10, c=RGBColor(0xCC, 0xCC, 0xCC))

    tx(s, Inches(1.2), Inches(5.5), Inches(3.5), Inches(0.3), "Certification", 10, True, accent)
    bl(s, Inches(1.2), Inches(5.9), Inches(3.5), Inches(0.5), [
        "Google Analytics (GA4) - 2026.04",
    ], sz=10, c=RGBColor(0xCC, 0xCC, 0xCC))

    # 우측 경력
    tx(s, Inches(5.5), Inches(0.5), Inches(7), Inches(0.4), "Experience", 20, True, BLK)

    # 에이원
    rrect(s, Inches(5.5), Inches(1.2), Inches(7.3), Inches(2.8))
    tx(s, Inches(5.8), Inches(1.3), Inches(5), Inches(0.4), "에이원퍼포먼스팩토리", 15, True, BLK)
    tx(s, Inches(5.8), Inches(1.7), Inches(5), Inches(0.3), "데이터인텔리전스팀 | 매니저  2025.03~2025.12 (10개월)", 10, c=GR)
    bl(s, Inches(5.8), Inches(2.1), Inches(6.8), Inches(1.7), [
        "광고 API 데이터 수집·가공 시스템 구축 (네이버/카카오)",
        "이용자 세그먼트 대량 변경 자동화 → 월 수작업 80% 절감",
        "광고 성과 리포트 자동화 → 리드타임 1일→실시간",
        "Validation 적용 → 입력 오류 장애 0건",
    ], sz=10)

    # 주네스
    rrect(s, Inches(5.5), Inches(4.2), Inches(7.3), Inches(1.8))
    tx(s, Inches(5.8), Inches(4.3), Inches(5), Inches(0.4), "주네스글로벌코리아", 15, True, BLK)
    tx(s, Inches(5.8), Inches(4.7), Inches(5), Inches(0.3), "재경부 | 주임  2018.12~2024.04 (5년 5개월)", 10, c=GR)
    bl(s, Inches(5.8), Inches(5.1), Inches(6.8), Inches(0.7), [
        "자금·결산 관리, 월 결산 처리 시간 30% 단축",
        "전사 10종 문서 페이퍼리스 전환 → 제출 누락 0건",
    ], sz=10)

    # 스킬
    tx(s, Inches(5.5), Inches(6.2), Inches(7), Inches(0.3), "Skills", 13, True, accent)
    if version == "data":
        tx(s, Inches(5.5), Inches(6.6), Inches(7), Inches(0.3),
           "Python  SQL  Pandas  NumPy  Streamlit  Plotly  Tableau  DuckDB  BigQuery  GA4  Docker  Git", 10, c=GR)
    elif version == "plan":
        tx(s, Inches(5.5), Inches(6.6), Inches(7), Inches(0.3),
           "서비스기획  KPI설계  수익화전략  요구사항분석  Python  SQL  Streamlit  FastAPI  GA4  Figma  Jira  Confluence", 10, c=GR)
    else:
        tx(s, Inches(5.5), Inches(6.6), Inches(7), Inches(0.3),
           "Python  JavaScript  FastAPI  Streamlit  Gemini API  Supabase  DuckDB  BigQuery  Docker  Render  pytest  Git", 10, c=GR)


# ═══ PROJECT SCREENSHOT ═══
def proj_shot(prs, accent, num, title, img_path, caption):
    s = prs.slides.add_slide(prs.slide_layouts[6]); sbg(s, DK)
    sidebar(s, accent)
    tx(s, Inches(0.8), Inches(0.3), Inches(1), Inches(0.4), f"0{num}", 28, True, accent)
    tx(s, Inches(1.8), Inches(0.35), Inches(10), Inches(0.4), title, 16, True, W)
    # 이미지 (슬라이드 안에 맞도록 크기 제한)
    if os.path.exists(img_path):
        rrect(s, Inches(2), Inches(1.1), Inches(9), Inches(5.2), fill=RGBColor(0x20, 0x2A, 0x44))
        im(s, img_path, Inches(2.2), Inches(1.3), w=Inches(8.6))
    tx(s, Inches(2), Inches(6.5), Inches(9), Inches(0.3), caption, 10, c=GR, a=PP_ALIGN.CENTER)


# ═══ PROJECT DETAIL ═══
def proj_detail(prs, accent, num, title, subtitle, left_title, left_items, right_title, right_items, tech, version):
    s = prs.slides.add_slide(prs.slide_layouts[6]); sbg(s, BG)
    sidebar(s, accent)

    # 헤더
    rect(s, Inches(0.25), Inches(0), Inches(13.1), Inches(0.9), fill=DK)
    tx(s, Inches(0.8), Inches(0.1), Inches(1), Inches(0.35), f"0{num}", 24, True, accent)
    tx(s, Inches(1.6), Inches(0.1), Inches(10), Inches(0.35), title, 16, True, W)
    tx(s, Inches(1.6), Inches(0.45), Inches(10), Inches(0.25), subtitle, 10, c=GR)

    # 좌측 카드
    rrect(s, Inches(0.8), Inches(1.2), Inches(5.8), Inches(4.2))
    rect(s, Inches(0.8), Inches(1.2), Inches(5.8), Inches(0.45), fill=accent)
    tx(s, Inches(1.1), Inches(1.23), Inches(5.3), Inches(0.35), left_title, 12, True, W)
    bl(s, Inches(1.1), Inches(1.8), Inches(5.3), Inches(3.3), left_items, sz=11)

    # 우측 카드
    rrect(s, Inches(7), Inches(1.2), Inches(5.8), Inches(4.2))
    rect(s, Inches(7), Inches(1.2), Inches(5.8), Inches(0.45), fill=GN)
    tx(s, Inches(7.3), Inches(1.23), Inches(5.3), Inches(0.35), right_title, 12, True, W)
    bl(s, Inches(7.3), Inches(1.8), Inches(5.3), Inches(3.3), right_items, sz=11)

    # 하단 Tech
    rrect(s, Inches(0.8), Inches(5.7), Inches(12), Inches(1.2))
    tx(s, Inches(1.1), Inches(5.8), Inches(2), Inches(0.3), "Tech Stack", 11, True, GR)
    tx(s, Inches(1.1), Inches(6.15), Inches(11.5), Inches(0.35), tech, 11, True, accent)


# ═══ ENDING ═══
def ending(prs, accent):
    s = prs.slides.add_slide(prs.slide_layouts[6]); sbg(s, DK)
    rect(s, 0, 0, Inches(0.4), Inches(7.5), fill=accent)
    tx(s, Inches(1.5), Inches(2), Inches(10), Inches(1), "Thank you.", 48, True, W)
    rect(s, Inches(1.5), Inches(3.3), Inches(3), Inches(0.01), fill=accent)
    tx(s, Inches(1.5), Inches(3.8), Inches(10), Inches(0.4), "이기쁨  Lee Gippeum", 18, True, W)
    tx(s, Inches(1.5), Inches(4.5), Inches(10), Inches(0.3), "ldsjoy@naver.com", 13, c=GR)
    tx(s, Inches(1.5), Inches(5), Inches(10), Inches(0.3), "github.com/rena-data", 13, c=accent)


# ═══ 프로젝트 데이터 ═══
PROJECTS = {
    "fintech": {
        "title": "Fintech Intelligence",
        "subtitle": "소규모 핀테크 경영 인텔리전스 시스템",
        "img": f"{SS}/fintech_dashboard.png",
        "caption": "Streamlit 4개 탭: 마케팅 성과 / 프로덕트 지표 / 경영 현황 / 시나리오 시뮬레이터",
        "tech": "Python  |  DuckDB  |  BigQuery  |  Pandas  |  Streamlit  |  Plotly  |  n8n  |  Docker",
        "data": {
            "left": ("데이터 분석 포인트", [
                "75,000+ 유저, 190만 이벤트 데이터 설계",
                "3 Layer 지표: 마케팅(CPC/CAC/ROAS) → 프로덕트(리텐션) → 재무(MRR)",
                "코호트 리텐션 히트맵 + 전환 퍼널 분석",
                "What-if 시나리오 시뮬레이터 (12개월)",
                "18개월 시계열 (시즌성, 주말효과 반영)",
            ]),
            "right": ("핵심 성과", [
                "4개 탭 인터랙티브 대시보드",
                "수동 작업 90% 절감 자동화 설계",
                "의사결정: '예산 어디에?', '런웨이 몇 개월?'",
                "DuckDB → BigQuery 무중단 전환 아키텍처",
            ]),
        },
        "plan": {
            "left": ("기획 & 전략", [
                "[Pain Point] 데이터 수집에 일일 1~2시간",
                "[타겟] 초기 핀테크 스타트업",
                "[솔루션] 자동 수집→대시보드→이상 감지→시뮬레이션",
                "[지표] 3 Layer: 마케팅→프로덕트→재무",
                "[자동화] n8n 워크플로우 4개 설계",
            ]),
            "right": ("핵심 성과", [
                "데이터 수집 일일 1~2시간 → 자동화 (90% 절감)",
                "리포트 수동 작성 → 자동 생성/배포",
                "이상 감지: 1일 → 10분 이내",
                "시나리오 분석: 건당 30분 → 실시간",
            ]),
        },
        "dev": {
            "left": ("아키텍처 & 기술", [
                "[DB] DuckDB ↔ BigQuery 환경변수 기반 전환",
                "[파이프라인] Python→데이터 생성→DuckDB→SQL 분석",
                "[대시보드] Streamlit 4탭 (Plotly 차트)",
                "[자동화] n8n 오케스트레이션 설계",
                "[스케일] 190만 이벤트 로컬 처리",
            ]),
            "right": ("기술적 도전 & 해결", [
                "동일 SQL이 DuckDB/BigQuery 모두 동작하도록 설계",
                "시즌성+주말효과+마케팅 상관관계 데이터 생성",
                "슬라이더→12개월 예측→실시간 차트 갱신",
                "환경변수 하나로 DB 무중단 전환",
            ]),
        },
    },
    "beauty": {
        "title": "Beauty AI Analyze",
        "subtitle": "AI 퍼스널컬러 분석 & 뷰티/패션 추천 (실서비스 운영)",
        "img": f"{SS}/beauty_main.png",
        "caption": "실서비스: beauty-ai-analyze.onrender.com  |  운영비 $0/월",
        "tech": "Python  |  FastAPI  |  Gemini Vision API  |  Supabase  |  JavaScript  |  Render  |  GA4",
        "data": {
            "left": ("데이터 분석 포인트", [
                "GA4 이벤트 7개 설계 (분석/클릭/다운로드/공유)",
                "Supabase: 분석 통계 + 제품 클릭 추적",
                "daily_stats: 일별 시즌 타입 분포",
                "popular_products: 인기 제품 TOP 50",
                "제품 큐레이션: 올리브영 80개 + 무신사 37개",
            ]),
            "right": ("핵심 성과", [
                "실서비스 배포 + 운영 ($0/월)",
                "4계절 매칭률 시각화 + S/A/B/C 등급",
                "레이더 차트 (밝기/채도/대비/온기/선명도)",
                "성별별 제품 추천 (남성/여성/유니섹스)",
            ]),
        },
        "plan": {
            "left": ("기획 & 전략", [
                "[시장] 컨설팅 1회 5~15만원, MZ 수요 폭발",
                "[전략] 무료 AI 진단→트래픽→제휴 수익화",
                "[MVP] 계획의 30%만 실행, $0 출시",
                "[차별화] 분석+추천+매칭+퀴즈+공유 올인원",
                "[수익화] 제휴 링크→포털 확장→프리미엄",
            ]),
            "right": ("핵심 성과", [
                "실서비스 배포: beauty-ai-analyze.onrender.com",
                "8개 탭 + SNS 공유 (X, Facebook)",
                "문의→Google Sheets + 이메일 알림 자동화",
                "개인정보처리방침 + 쿠키 동의 + GA4",
            ]),
        },
        "dev": {
            "left": ("아키텍처 & 트러블슈팅", [
                "[Backend] FastAPI + Gemini Vision API",
                "[AI] 모델 폴백 4개 + 키 이중화 → 일일 3,000회",
                "[DB] Supabase - 통계/추적/공유URL",
                "[503] 2.5-flash→lite→2.0→3.0 순차 시도",
                "[Streamlit→JS] HTML 깨짐 → 완전 전환",
            ]),
            "right": ("기술적 도전 & 해결", [
                "temperature 0.0 → 동일 사진 동일 결과",
                "html2canvas span 크기0 → div+인라인 크기",
                "FormData season_type → Form() 파라미터",
                "SVG 레이더 차트 라이브러리 없이 구현",
            ]),
        },
    },
    "pet": {
        "title": "Raising Pet",
        "subtitle": "킹받는 펫 키우기 게임 (상태 머신 기반)",
        "img": f"{SS}/pet_main.png",
        "caption": "5종 동물 × 4성장단계 × 3난이도  |  26개 단위 테스트 100% 통과",
        "tech": "Python  |  JavaScript  |  Streamlit  |  SVG  |  CSS Animation  |  localStorage  |  pytest",
        "data": {
            "left": ("데이터 분석 포인트", [
                "GA4 이벤트 12개 설계 (게임/액션/성장/공유)",
                "유저 퍼널: 인지→유입→활성화→리텐션",
                "밸런스: 3스탯 × 4단계 × 3난이도 매트릭스",
                "수익 시뮬레이션: eCPM × DAU 계산",
                "리텐션 KPI: D1 30%, D7 15% 목표",
            ]),
            "right": ("핵심 성과", [
                "복합 밸런스: 연쇄 붕괴 + 피로도 + 랜덤 이벤트",
                "DAU 1K→월 15~45만원 수익 시뮬레이션",
                "4단계 그로스 전략 ($0→바이럴→광고)",
                "6가지 배포 옵션 비교 분석",
            ]),
        },
        "plan": {
            "left": ("기획 & 바이럴 전략", [
                "[컨셉] '결국 죽는'→감정 유발→자발적 공유",
                "[바이럴] 추모카드 + X 공유 + 챌린지",
                "[게이미피케이션] 칭호 (학대범→전설의 집사)",
                "[그로스] $0→바이럴→챌린지→광고→리텐션",
                "[수익] DAU 10K→월 150~450만원",
            ]),
            "right": ("핵심 성과", [
                "MVP 완성 + 배포 준비 완료",
                "X 자동 공유 + 추모카드 바이럴 설계",
                "칭호 시스템으로 경쟁 유도",
                "6가지 배포 옵션 + 로드맵 수립",
            ]),
        },
        "dev": {
            "left": ("아키텍처 & 기술", [
                "[패턴] 상태 머신 (ALIVE↔DEAD)",
                "[Frontend] 순수 JS + SVG 렌더러 (20변형)",
                "[통신] postMessage 양방향 (Python↔JS)",
                "[밸런스] 연쇄 붕괴 + 피로도 + 확률 이벤트",
                "[테스트] pytest 26개 100% 통과",
            ]),
            "right": ("기술적 도전 & 해결", [
                "성장 단계별 감소 배율 0.8x~1.6x 설계",
                "액션 피로도: 연속 시 100%→70%→40%",
                "구버전 세이브 호환성 100% 직렬화",
                "SVG 5종 × 4기분 = 20가지 동적 렌더링",
            ]),
        },
    },
}


def generate(version, accent):
    labels = {"data": "데이터 분석", "plan": "서비스 기획", "dev": "개발"}
    subs = {
        "data": "데이터 수집 → 분석 → 시각화 → 인사이트 → 의사결정 지원",
        "plan": "문제 정의 → 서비스 설계 → MVP → 데이터 기반 개선 → 수익화",
        "dev": "설계 → 구현 → 테스트 → 배포 → 운영",
    }

    prs = nprs()
    cover(prs, f"{labels[version]} 포트폴리오", subs[version], accent)
    about(prs, accent, version)

    for i, key in enumerate(["fintech", "beauty", "pet"], 1):
        p = PROJECTS[key]
        proj_shot(prs, accent, i, f"{p['title']} — {p['subtitle']}", p["img"], p["caption"])
        d = p[version]
        proj_detail(prs, accent, i, p["title"], p["subtitle"], d["left"][0], d["left"][1], d["right"][0], d["right"][1], p["tech"], version)

    ending(prs, accent)
    return prs


# ─── Generate ───
accents = {"data": GN, "plan": RGBColor(0x7C, 0x4D, 0xFF), "dev": BL}
names = {"data": "DataAnalyst", "plan": "Planner", "dev": "Developer"}

for v in ["data", "plan", "dev"]:
    prs = generate(v, accents[v])
    prs.save(f"{OUT}/Portfolio_{names[v]}.pptx")
    print(f"Portfolio_{names[v]}.pptx 생성 완료")

print("\n3개 포트폴리오 모두 완료!")
